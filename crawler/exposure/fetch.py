"""
Fetch and normalise candidate documents for COLLECT-007.

Downloads a URL and converts it to clean text: pymupdf for PDF, python-docx
for DOCX, and simple tag-stripping for HTML. Each normalised document keeps
per-section location markers (page number, paragraph index) for provenance.
"""

from __future__ import annotations

import io
import logging
import re
from typing import Literal

import httpx
from pydantic import BaseModel, Field

from crawler.exposure.discovery import USER_AGENT
from crawler.exposure.domains import DomainMention, dedupe, find_domain_mentions, mention_from_link_uri

log = logging.getLogger(__name__)

DocumentType = Literal["pdf", "docx", "html"]

_SCRIPT_STYLE = re.compile(r"<(script|style)[^>]*>.*?</\1>", re.IGNORECASE | re.DOTALL)
_TAG = re.compile(r"<[^>]+>")
_WHITESPACE = re.compile(r"[ \t]+")
_BLANK_LINES = re.compile(r"\n\s*\n+")
# Only <a href>: excludes <link href> (stylesheet/feed/oembed/RSD discovery
# tags in <head>) and <script src>, which are template plumbing, not
# content a visitor would actually see or click.
_ANCHOR_HREF = re.compile(r'<a\b[^>]*?\bhref\s*=\s*["\']([^"\']+)["\']', re.IGNORECASE)


class FetchError(RuntimeError):
    pass


class RateLimitedError(FetchError):
    """The target returned HTTP 429. Carries Retry-After (seconds) when the site sent one."""

    def __init__(self, url: str, retry_after: float | None) -> None:
        super().__init__(f"{url}: HTTP 429 (rate-limited)")
        self.retry_after = retry_after


class Section(BaseModel):
    location: str
    text: str


class NormalizedDocument(BaseModel):
    url: str
    document_type: DocumentType
    sections: list[Section]
    domain_mentions: list[DomainMention] = Field(default_factory=list)


class FetchHttpClient:
    def __init__(self, http: httpx.Client | None = None) -> None:
        self._http = http or httpx.Client(
            timeout=30.0, follow_redirects=True, headers={"User-Agent": USER_AGENT}
        )

    def get(self, url: str) -> tuple[bytes, str]:
        try:
            response = self._http.get(url)
        except httpx.HTTPError as exc:
            raise FetchError(f"{url}: {type(exc).__name__}") from None
        if response.status_code == 429:
            retry_after = response.headers.get("retry-after")
            raise RateLimitedError(url, float(retry_after) if retry_after else None)
        if response.status_code >= 400:
            raise FetchError(f"{url}: HTTP {response.status_code}")
        return response.content, response.headers.get("content-type", "")


def detect_document_type(url: str, content_type: str) -> DocumentType:
    lowered_url = url.lower()
    lowered_type = content_type.lower()
    if lowered_url.endswith(".pdf") or "application/pdf" in lowered_type:
        return "pdf"
    if lowered_url.endswith(".docx") or "officedocument.wordprocessingml" in lowered_type:
        return "docx"
    return "html"


def normalize_pdf(content: bytes, url: str) -> NormalizedDocument:
    import pymupdf

    sections: list[Section] = []
    mentions: list[DomainMention] = []
    with pymupdf.open(stream=content, filetype="pdf") as doc:
        for page_number, page in enumerate(doc, start=1):
            location = f"page {page_number}"
            text = page.get_text().strip()
            if text:
                sections.append(Section(location=location, text=text))
                mentions.extend(find_domain_mentions(text, source_url=url, location=location))
            for link in page.get_links():
                uri = link.get("uri")
                if not uri:
                    continue
                mention = mention_from_link_uri(uri, source_url=url, location=f"{location} (link)")
                if mention:
                    mentions.append(mention)
    return NormalizedDocument(url=url, document_type="pdf", sections=sections, domain_mentions=mentions)


def normalize_docx(content: bytes, url: str) -> NormalizedDocument:
    import docx

    document = docx.Document(io.BytesIO(content))
    sections: list[Section] = []
    mentions: list[DomainMention] = []
    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = paragraph.text.strip()
        if text:
            location = f"paragraph {index}"
            sections.append(Section(location=location, text=text))
            mentions.extend(find_domain_mentions(text, source_url=url, location=location))
    return NormalizedDocument(url=url, document_type="docx", sections=sections, domain_mentions=mentions)


def normalize_html(content: bytes, url: str) -> NormalizedDocument:
    raw = content.decode("utf-8", errors="replace")

    without_scripts = _SCRIPT_STYLE.sub(" ", raw)
    text = _TAG.sub(" ", without_scripts)
    text = _WHITESPACE.sub(" ", text)
    text = _BLANK_LINES.sub("\n", text).strip()

    # Visible text for emails/URLs/www-mentions, plus actual <a href> link
    # targets (which may not be shown as text, e.g. mailto: links). This
    # deliberately skips <script src>, <link>, <img>, and other markup
    # plumbing, which would otherwise flood the report with the site's own
    # CSS/JS/feed URLs rather than meaningful disclosures.
    mentions = find_domain_mentions(text, source_url=url, location="page")
    for href in _ANCHOR_HREF.findall(raw):
        mention = mention_from_link_uri(href, source_url=url, location="page (link)")
        if mention:
            mentions.append(mention)

    return NormalizedDocument(
        url=url,
        document_type="html",
        sections=[Section(location="page", text=text)] if text else [],
        domain_mentions=dedupe(mentions),
    )


_NORMALIZERS = {"pdf": normalize_pdf, "docx": normalize_docx, "html": normalize_html}


class DocumentFetcher:
    def __init__(self, http: FetchHttpClient | None = None) -> None:
        self.http = http or FetchHttpClient()

    def fetch(self, url: str) -> NormalizedDocument:
        content, content_type = self.http.get(url)
        document_type = detect_document_type(url, content_type)
        try:
            return _NORMALIZERS[document_type](content, url)
        except Exception as exc:  # normalisation failures shouldn't abort the whole run
            raise FetchError(f"{url}: could not normalise as {document_type} ({exc})") from exc
