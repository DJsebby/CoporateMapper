"""Offline tests for crawler.exposure.fetch. No network calls are made."""

import io
import unittest

import httpx

from crawler.exposure.fetch import (
    DocumentFetcher,
    FetchError,
    FetchHttpClient,
    RateLimitedError,
    detect_document_type,
    normalize_docx,
    normalize_html,
    normalize_pdf,
)


class FakeHttp:
    def __init__(self, content: bytes, content_type: str) -> None:
        self.content = content
        self.content_type = content_type

    def get(self, url: str):
        return self.content, self.content_type


class DetectDocumentTypeTests(unittest.TestCase):
    def test_pdf_by_extension(self):
        self.assertEqual(detect_document_type("https://acme.example/x.pdf", ""), "pdf")

    def test_pdf_by_content_type(self):
        self.assertEqual(
            detect_document_type("https://acme.example/x", "application/pdf"), "pdf"
        )

    def test_docx_by_extension(self):
        self.assertEqual(detect_document_type("https://acme.example/x.docx", ""), "docx")

    def test_html_default(self):
        self.assertEqual(detect_document_type("https://acme.example/careers", "text/html"), "html")


class NormalizeHtmlTests(unittest.TestCase):
    def test_strips_tags_and_scripts(self):
        html = b"""
        <html><head><style>.a{color:red}</style></head>
        <body><script>evil()</script><h1>Hello</h1><p>World policy text</p></body></html>
        """
        doc = normalize_html(html, "https://acme.example/policy")
        self.assertEqual(len(doc.sections), 1)
        self.assertNotIn("evil", doc.sections[0].text)
        self.assertIn("Hello", doc.sections[0].text)
        self.assertIn("World policy text", doc.sections[0].text)

    def test_empty_body_produces_no_sections(self):
        doc = normalize_html(b"<html><body></body></html>", "https://acme.example/empty")
        self.assertEqual(doc.sections, [])

    def test_mailto_href_captured_even_though_stripped_from_visible_text(self):
        html = b'<html><body><a href="mailto:jane@acme.example">Email us</a></body></html>'
        doc = normalize_html(html, "https://acme.example/contact")
        self.assertNotIn("jane@acme.example", doc.sections[0].text)
        [mention] = doc.domain_mentions
        self.assertEqual(mention.kind, "email")
        self.assertEqual(mention.domain, "acme.example")

    def test_visible_url_captured(self):
        html = b"<html><body><p>Read https://vendor.example/terms</p></body></html>"
        doc = normalize_html(html, "https://acme.example/terms")
        [mention] = doc.domain_mentions
        self.assertEqual(mention.domain, "vendor.example")

    def test_retina_image_srcset_is_not_a_false_positive_email(self):
        html = (
            b'<html><body><img srcset="logo-default@2x-576x432.png 2x, '
            b'logo-default@1x-288x216.png 1x"></body></html>'
        )
        doc = normalize_html(html, "https://acme.example/")
        self.assertEqual(doc.domain_mentions, [])

    def test_stylesheet_script_and_head_metadata_links_are_not_scanned(self):
        html = (
            b'<html><head>'
            b'<link rel="stylesheet" href="https://acme.example/wp-includes/css/dashicons.min.css?ver=7.1.2">'
            b'<link rel="alternate" type="application/rss+xml" href="https://acme.example/feed/">'
            b'<script src="https://acme.example/wp-includes/js/jquery.js"></script>'
            b'</head><body>'
            b'<a href="https://vendor.example/portal">Vendor portal</a>'
            b'</body></html>'
        )
        doc = normalize_html(html, "https://acme.example/about")
        domains = {m.domain for m in doc.domain_mentions}
        self.assertEqual(domains, {"vendor.example"})

    def test_anchor_link_to_own_domain_still_captured(self):
        html = b'<html><body><a href="https://acme.example/careers">Careers</a></body></html>'
        doc = normalize_html(html, "https://acme.example/about")
        domains = {m.domain for m in doc.domain_mentions}
        self.assertEqual(domains, {"acme.example"})


class NormalizePdfTests(unittest.TestCase):
    def test_pages_extracted_with_location(self):
        import pymupdf

        pdf = pymupdf.open()
        page = pdf.new_page()
        page.insert_text((72, 72), "All staff work fully remote from home offices.")
        content = pdf.tobytes()
        pdf.close()

        doc = normalize_pdf(content, "https://acme.example/policy.pdf")
        self.assertEqual(len(doc.sections), 1)
        self.assertEqual(doc.sections[0].location, "page 1")
        self.assertIn("remote", doc.sections[0].text)

    def test_visible_email_captured_as_domain_mention(self):
        import pymupdf

        pdf = pymupdf.open()
        page = pdf.new_page()
        page.insert_text((72, 72), "Contact info@acme.example for details.")
        content = pdf.tobytes()
        pdf.close()

        doc = normalize_pdf(content, "https://acme.example/policy.pdf")
        self.assertIn("acme.example", [m.domain for m in doc.domain_mentions])

    def test_link_annotation_captured_even_without_visible_text(self):
        import pymupdf

        pdf = pymupdf.open()
        page = pdf.new_page()
        page.insert_text((72, 72), "Click here.")
        page.insert_link(
            {"kind": pymupdf.LINK_URI, "from": pymupdf.Rect(72, 72, 150, 90), "uri": "https://vendor.example/portal"}
        )
        content = pdf.tobytes()
        pdf.close()

        doc = normalize_pdf(content, "https://acme.example/policy.pdf")
        [mention] = [m for m in doc.domain_mentions if m.kind == "link"]
        self.assertEqual(mention.domain, "vendor.example")


class NormalizeDocxTests(unittest.TestCase):
    def test_paragraphs_extracted_with_location(self):
        import docx

        document = docx.Document()
        document.add_paragraph("Devices are company-managed laptops only.")
        buffer = io.BytesIO()
        document.save(buffer)

        doc = normalize_docx(buffer.getvalue(), "https://acme.example/policy.docx")
        self.assertEqual(len(doc.sections), 1)
        self.assertEqual(doc.sections[0].location, "paragraph 1")
        self.assertIn("company-managed", doc.sections[0].text)

    def test_visible_email_captured_as_domain_mention(self):
        import docx

        document = docx.Document()
        document.add_paragraph("Questions? Email hr@acme.example.")
        buffer = io.BytesIO()
        document.save(buffer)

        doc = normalize_docx(buffer.getvalue(), "https://acme.example/policy.docx")
        self.assertIn("acme.example", [m.domain for m in doc.domain_mentions])


class FetchHttpClientRateLimitTests(unittest.TestCase):
    def _client_for(self, response: httpx.Response) -> FetchHttpClient:
        transport = httpx.MockTransport(lambda request: response)
        return FetchHttpClient(http=httpx.Client(transport=transport))

    def test_429_with_retry_after_raises_with_seconds(self):
        client = self._client_for(httpx.Response(429, headers={"Retry-After": "7"}))
        with self.assertRaises(RateLimitedError) as ctx:
            client.get("https://acme.example/throttled")
        self.assertEqual(ctx.exception.retry_after, 7.0)

    def test_429_without_retry_after_raises_with_none(self):
        client = self._client_for(httpx.Response(429))
        with self.assertRaises(RateLimitedError) as ctx:
            client.get("https://acme.example/throttled")
        self.assertIsNone(ctx.exception.retry_after)


class DocumentFetcherTests(unittest.TestCase):
    def test_fetch_dispatches_html_normalizer(self):
        fetcher = DocumentFetcher(http=FakeHttp(b"<html><body>Hi</body></html>", "text/html"))
        doc = fetcher.fetch("https://acme.example/careers")
        self.assertEqual(doc.document_type, "html")
        self.assertIn("Hi", doc.sections[0].text)

    def test_normalisation_failure_raises_fetch_error(self):
        fetcher = DocumentFetcher(http=FakeHttp(b"not a pdf", "application/pdf"))
        with self.assertRaises(FetchError):
            fetcher.fetch("https://acme.example/broken.pdf")


if __name__ == "__main__":
    unittest.main()
